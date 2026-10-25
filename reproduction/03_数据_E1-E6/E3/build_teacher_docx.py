"""Create a teacher-facing explanation of archived E3 results without rescoring."""
from pathlib import Path
import hashlib
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[1]
SOURCE = ROOT / 'docs/results/2026-10-07/e3-b5ad6d180c9f42ce/FINAL-RESULTS.json'
raw = SOURCE.read_bytes()
data = json.loads(raw)
assert data['collection']['actual'] == 216
assert data['rating']['actual'] == 216
assert data['expert']['status'] == 'pending'
OUT = BASE / 'output/documents/E3-实验结果说明-教师审阅稿.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.65), Cm(1.65)
sec.left_margin, sec.right_margin = Cm(2), Cm(2)
sec.footer_distance = Cm(.75)

def font(style, size, bold=False):
    style.font.name = 'Arial Unicode MS'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(fonts.attrib):
        del fonts.attrib[key]
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + key), 'Arial Unicode MS')
    for element in style.element.xpath('./w:pPr/w:pBdr | ./w:pPr/w:numPr'):
        element.getparent().remove(element)
    style.font.italic = False

font(doc.styles['Normal'], 10.5)
n = doc.styles['Normal'].paragraph_format
n.line_spacing = 1.12
n.space_after = Pt(4)
font(doc.styles['Title'], 21, True)
doc.styles['Title'].paragraph_format.space_after = Pt(5)
font(doc.styles['Subtitle'], 10)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(10)
font(doc.styles['Heading 1'], 12, True)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(5)

def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text):
    p(text, 'Heading 1')

def table(headers, rows, widths, numeric=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        for k, v in {'val':'single', 'sz':'5', 'color':'D9D9D9'}.items():
            el.set(qn('w:' + k), v)
        borders.append(el)
    pr.append(borders)
    for i, row in enumerate([headers] + rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        trpr = cells[0]._tc.getparent().get_or_add_trPr()
        avoid = OxmlElement('w:cantSplit')
        trpr.append(avoid)
        if i == 0:
            repeat = OxmlElement('w:tblHeader')
            trpr.append(repeat)
        for j, (cell, text, width) in enumerate(zip(cells, row, widths)):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge, value in [('top',85),('bottom',85),('left',100),('right',100)]:
                e = OxmlElement('w:' + edge)
                e.set(qn('w:w'), str(value))
                e.set(qn('w:type'), 'dxa')
                margins.append(e)
            cp.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '23435C' if i == 0 else ('F3F6F8' if i % 2 else 'FFFFFF'))
            cp.append(shade)
            pp = cell.paragraphs[0]
            pp.paragraph_format.space_after = Pt(0)
            pp.paragraph_format.line_spacing = 1.12
            if j in numeric:
                pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = pp.add_run(text)
            rr.font.size = Pt(9.5)
            if i == 0:
                rr.bold = True
                rr.font.color.rgb = RGBColor(255,255,255)
    return t

def metric(group, key):
    return f"{data['main']['systems'][group]['metrics'][key]['median']:.2%}"

p('实验三科学语义评测结果说明', 'Title')
p('教师审阅稿　2026年10月8日', 'Subtitle')
p('当前结果：完整领域材料组的关键科学断言命中率为90.14%，普通接口说明组为66.20%，相差23.94个百分点。实验已完成回答采集和助手主评，独立天文专家复核尚未完成。')
h('一　实验目的与设置')
p('本实验检验：给模型提供完整的天文领域材料，能否减少科学语义错误与无依据推断。两组使用相同的deepseek-flash模型和生成参数；本轮结果不能归属于GPT-6.1。')
table(['组别', '向被测模型提供的材料'], [
    ['A组（S-A）', '完整skill及天文领域材料，包括科学概念、数据列含义和能力边界。'],
    ['B组（S-B）', '普通API包装器，仅提供表名、列名和两阶段查询协议。'],
], [3.1,13.9])
p('题集包含32道离线主测题、4道在线辅测题，每题两组各独立回答3次，共216份真实回答：主测192份，辅测24份。以下主结果仅统计离线题，在线题另行分析。')
h('二　主要结果及读法')
p('下表报告三次重复结果的中位数，即将三个数排序后居中的数。各项指标的分母不同，不能相加为总分。')
rows = []
for key, name, meaning in [
    ('sfa','关键科学断言命中率 SFA','应覆盖的关键科学点，实际覆盖多少；越高越好。'),
    ('ungr','无来源数值率 UNGR','所列科学数字中，缺少可追溯来源的比例；越低越好。'),
    ('ucr','无支撑主张率 UCR','所列科学说法中，缺少充分依据的比例；越低越好。'),
    ('slr','静态误称时序率 SLR','相关题中，将汇总测光误作时间变化证据的比例；越低越好。'),
    ('rqc','必要限定语覆盖率 RQC','应交代的来源、范围及不确定性等条件的覆盖比例；越高越好。'),
]:
    rows.append([name, meaning, metric('S-A',key), metric('S-B',key)])
table(['指标','含义与判断方向','A组','B组'], rows, [4.3,8.5,2.1,2.1], (2,3))
p('90.14%不是“90.14%的题完全答对”。每次主测有71个关键科学点；A组其中两次命中64个，64÷71＝90.14%。一题可能包含多个科学点。限定语每次共73项；时序相关题每次共14道。')
p(f'按更严格的逐回答合格规则，两组合格率中位数分别为{metric("S-A","pass_rate")}和{metric("S-B","pass_rate")}。第六项人机评分一致性因缺少专家独立评分，暂不能计算。')

doc.add_page_break()
h('三　实际错误与结果含义')
p('以下为现有回答中的代表性片段及当前主评判断，仍待专家复核。')
p('① 将汇总测光当成光变证据。B组称“在 Stripe82 观测时段内它在 r 波段是稳定的”。汇总星等和散射不等于完整光变曲线，不能据此确认整个观测时段物理恒定。')
p('② 条件不足仍给出精确数值。B组给出窄波段星等“15.44”，但题面缺少完整单位、带通响应及零点，不能将公式计算值当成已核实的实际星等。')
p('③ 把缺数据读成无变化。B组根据has_epoch_photometry=False，推断“不会有光变污染问题”。没有公开历元测光只说明数据可用性，不能证明该源没有光变。')
p('当前主评在B组16份回答中检出科学红线；A组未检出红线，但仍有科学点遗漏与无支撑说法。三份已知坏回答均被识别，只说明反证检查通过。')
h('四　结果文件中几个容易误解的字段')
table(['字段', '实际含义'], [
    ['collection.actual = 216', '216份回答采集完成；不代表回答全部正确。'],
    ['rating.valid = true', '评分记录及反证检查符合流程要求；不代表科学表现全部合格。'],
    ['scientific_performance.status = FAIL', '存在触发科学红线的回答；不表示实验未运行或所有回答都错。'],
    ['median / per_repeat / range', '三次中位数／三次各自结果／最小值与最大值；range不是置信区间。'],
    ['expert.status = pending；agreement = null', '专家复核未完成，因此人机一致性暂不可计算；null不等于零分。'],
], [7.0,10.0])
h('五　结论边界与待复核事项')
p('评分由当前对话助手直接完成，未调用外部评分API。助手参与准备并见过组别映射，不构成独立盲评；需一名未参与出题的天文专家独立评分，并核对科学主张及数值的枚举分母。“无来源”不等于数值必然错误。')
p('材料也存在待复核问题：E04题面假定表中有Teff/logg，但公开列清单缺少这两个输出列；后续E6核验记录还发现，部分材料将Download栏的文件大小误写成星数。相应题目、依据和失分需要复核，本说明未事后改分。')
p('统计补充：按题号成组重抽10,000次，关键断言命中率差的95%区间为12.68至33.33个百分点；逐题三次多数合格的精确McNemar检验p＝0.00048828。统计依赖当前评分，不能消除评分偏差。只有32个独立主测题号；两组信息量和上下文长度也不同，不能据此单独归因于skill架构。')
p('建议结论表述：在本题集与当前助手评分下，完整领域材料组整体表现较好；独立专家复核尚未完成，不能宣称已消除幻觉或达到专家水平。')
last = p('数据依据：批次e3-b5ad6d180c9f42ce的FINAL-RESULTS.json及FINAL-REPORT.md；材料补充核验见2026年10月7日E6最终报告。本文件仅解释既有结果，未重新运行实验或修改评分。')
for run in last.runs:
    run.font.size = Pt(8)

foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = foot.add_run('第 ')
r.font.size = Pt(8)
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
foot._p.append(fld)
foot.add_run(' 页').font.size = Pt(8)
doc.core_properties.title = '实验三科学语义评测结果说明'
doc.core_properties.subject = '教师审阅稿'
doc.core_properties.author = ''
for paragraph in list(doc.paragraphs) + [pp for tt in doc.tables for rr in tt.rows for cc in rr.cells for pp in cc.paragraphs] + list(sec.footer.paragraphs):
    for run in paragraph.runs:
        fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        for key in list(fonts.attrib):
            del fonts.attrib[key]
        for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + key), 'Arial Unicode MS')
doc.save(OUT)
assert SOURCE.read_bytes() == raw
print(json.dumps({'output':str(OUT), 'source_unchanged':True, 'sha256':hashlib.sha256(raw).hexdigest()}, ensure_ascii=False))
