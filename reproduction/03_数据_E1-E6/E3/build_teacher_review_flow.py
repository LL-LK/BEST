"""根据已冻结的E3文件编写教师复核流程，不运行实验或修改评分。"""
from pathlib import Path
import ast
import hashlib
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
BATCH = BASE / 'batches/e3-b5ad6d180c9f42ce'
SOURCE = BATCH / 'FINAL-RESULTS.json'
original = SOURCE.read_bytes()
data = json.loads(original)
for name in ['bundle.json', 'PROTOCOL.md', 'answers/S-B_E3-A05_r1.json', 'scores/S-B_E3-A05_r1.json']:
    assert (BATCH / name).is_file()
OUT = BASE / 'output/documents/E3-教师复核流程.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(1.6)
sec.left_margin = sec.right_margin = Cm(2)
sec.footer_distance = Cm(.7)
# 仅复用已验证的排版函数，不执行旧报告生成程序。
tree = ast.parse((BASE / 'build_teacher_docx.py').read_text())
helpers = ast.Module(body=[node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name in {'font','p','h','table'}], type_ignores=[])
exec(compile(helpers, '<validated-document-formatting>', 'exec'), globals())
font(doc.styles['Normal'], 10.5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.1
doc.styles['Normal'].paragraph_format.space_after = Pt(4)
font(doc.styles['Title'], 19, True)
doc.styles['Title'].paragraph_format.space_after = Pt(4)
font(doc.styles['Subtitle'], 9)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(8)
font(doc.styles['Heading 1'], 12, True)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(5)

p('实验三结果文件教师复核流程', 'Title')
p('适用于既有结果的学术复核　2026年10月8日', 'Subtitle')
p('本流程适用于看过现有结果后的复核；未接触组别结果的独立盲评，应另用匿名包，不先阅读本流程中的已知评分示例。')
p('请您核查原始回答是否支持现有评分，以及汇总指标能否从逐项记录重算。FINAL-RESULTS.json是汇总与错误索引，仅凭这一个文件不能完成科学复核。下面按“准备材料→检查完整性→逐份核对→重算指标→反馈分歧”的顺序进行。')
h('第一步　准备同一批次的配套文件')
p('本次批次号为e3-b5ad6d180c9f42ce。请实验提供方将下列文件按原目录关系交付；不要混用其他批次。老师可使用带JSON格式化功能的本地文本编辑器打开文件，无需调用模型API。')
table(['文件或目录', '复核用途'], [
    ['FINAL-RESULTS.json', '查看完成状态、分组指标及error_taxonomy错误索引。'],
    ['FINAL-REPORT.md、PROTOCOL.md', '查看实验设计、指标定义、合格判据和已披露限制。'],
    ['bundle.json', '按items中的id定位题面与评分依据；同时核查依据本身。'],
    ['answers/、scores/', '查看真实回答、实际输入、逐条判定及计数；同名文件相互对应。'],
    ['live/、failures/、manifest.json', '核查在线原始回包、失败记录、冻结配置与文件哈希。'],
    ['材料补充核验记录', '同时提供归档E3报告及E6报告中关于题面、列清单和文件大小的核验说明。'],
], [6.1,10.9])
h('第二步　检查完整性和状态含义')
p('在汇总JSON中搜索batch_id、collection、rating、expert和delivery。应核对：回答216份＝离线192份＋在线24份；评分216份；被测模型为deepseek-flash。再检查每题每组是否都有r1、r2、r3，是否有重复、缺失或截断。数量相等仍需核对文件身份。')
p('rating.valid=true表示评分记录通过流程检查；scientific_performance.status=FAIL表示存在触发科学红线的回答；expert.status=pending与agreement=null表示专家复核未完成。以上字段均不能直接解释为“实验全部正确”或“实验没有运行”。')
doc.add_page_break()
h('第三步　从一条错误记录追溯到原文')
p('先在error_taxonomy中选一条记录，读出item_id、group、repeat、quote、reason和redline。例如E3-A05、S-B、第1次，对应answers/S-B_E3-A05_r1.json与scores/S-B_E3-A05_r1.json。')
p('先读bundle.json中的题面和输入数据，再读回答文件的answer及实际请求。最后读评分文件score内的critical、qualifiers、numbers、claims、forbidden、redlines，核对每个quote是否逐字存在、理由是否符合上下文、source能否定位到具体列或资料。')
p('该例红线片段为“根本没有 1.2 mag 这种量级的变化”。请独立判断：仅凭Stripe82汇总统计，是否足以作出完整历史振幅的结论。既要检查这条扣分是否合理，也要检查整篇回答有无漏计的数值、主张或限定语。不要只依据摘要同意原判。')
p('请按同样方法核查全部216份回答，既查扣分也查得分；error_taxonomy只列部分红线，不能代替全量复核。若只抽查，请记录抽样方法、题号和数量，结论限定为抽查结果。')

h('第四步　核对计数并重算五项主指标')
p('仅使用32道离线题，按组别和重复序号分别汇总scores中的counts。JSON路径main.systems.S-A.per_repeat及S-B对应三次计数；数组第1、2、3项分别为第1、2、3次。')
table(['指标', '每次重复的重算方法'], [
    ['SFA　关键科学断言命中率', 'Σcritical_hit ÷ Σcritical_total。'],
    ['UNGR　无来源数值率', 'Σunsourced_num ÷ Σnumbers_total。'],
    ['UCR　无支撑主张率', 'Σunsupported_claim ÷ Σclaims_total。'],
    ['SLR　静态误称时序率', 'bundle中slr_ids对应的主测题中，counts.static_as_lc为1的题数 ÷ 14。'],
    ['RQC　必要限定语覆盖率', 'Σqualifier_hit ÷ Σqualifier_total。'],
], [6.1,10.9])
p('核算例：A组第1次SFA＝65÷71＝91.55%，第2、3次均为64÷71＝90.14%；中位数为90.14%，range为[90.14%, 91.55%]。先汇总分子分母再相除，不把各题百分比简单平均；range是最小值与最大值，不是置信区间。分母为零时记缺测，不填0或1。')
p('严格合格率需另核counts.pass：关键断言全部命中、无禁止结论、无红线且无致命编数。SFA高不等于整份回答合格。在线24份仅核live部分，不并入main。')
doc.add_page_break()
h('第五步　优先复核这些争议与统计口径')
p('数值和主张是否完整枚举，不能只检查已列出的条目；无来源与数值错误应区分。来源可追溯也不保证来源本身正确。静态测光不统一等同于单次曝光；没有历元测光、未标记变星均不能证明物理恒定。')
p('请重点核查E04题面所需Teff/logg与公开输出列不一致，以及后续E6记录发现的“Download文件大小被误写为星数”。应分别记录题目或材料缺陷与模型错误，提出修订建议，保留原评分，不事后无说明地删题或改分。')
p('如复核统计显著性：独立题号应为32；McNemar按每题三次中至少两次合格作配对判定，当前仅A合格12题、仅B合格0题；bootstrap为按题号成组重抽10,000次，同时保留两组和三次重复，不能将192份回答当成192个独立样本。')
h('第六步　独立保存意见并交回')
p('请保留原FINAL-RESULTS.json，在单独文件记录每项分歧。建议每条包含：题号、组别、重复序号、评分字段及条目ID、原判、复核判定、回答原文、来源定位、理由、建议处理。结论可写“认可／需修订／证据不足”；证据不足不得强行记为通过。')
p('交回内容：复核范围与方法、逐项评分或分歧清单、需修订的题目或材料、可支持的结论。由实验提供方保存专家原件、核验完整性、生成修订版本与分歧报告，原助手评分不被自动覆盖。')
h('如需作为独立盲评专家')
p('应先只接收E3-expert-blind-219.zip，按README.md及index.csv顺序填写scorecard.blank.json的副本，使用samples和truth-reference等匿名依据，评分完成后再看本JSON及组别结果。met与triggered填写true/false，usability填1–5，理由至少8字，quote来自回答原文，数值及主张须全量枚举。')
p('匿名包共219项，含质量控制项，须全量独立完成。若此前已看过组别结果或助手评分，请如实记录已知信息，此次应称结果复核，不能称未受结果影响的独立盲评。抽查意见也不能代替全量人机一致性数据。')
last=p('说明：本流程依据该批次既有文件编写，仅提供复核方法，未代替老师评分，未改动实验结果。')
last.runs[0].font.size=Pt(8)

for paragraph in list(doc.paragraphs) + [pp for tt in doc.tables for rr in tt.rows for cc in rr.cells for pp in cc.paragraphs]:
    for run in paragraph.runs:
        fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        for key in list(fonts.attrib):
            del fonts.attrib[key]
        for key in ('ascii','hAnsi','eastAsia','cs'):
            fonts.set(qn('w:'+key),'Arial Unicode MS')
foot=sec.footer.paragraphs[0]
foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('第 ').font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
foot.add_run(' 页').font.size=Pt(8)
doc.core_properties.title='实验三结果文件教师复核流程'
doc.core_properties.author=''
doc.save(OUT)
assert SOURCE.read_bytes()==original
print(json.dumps({'output':str(OUT),'source_unchanged':True,'sha256':hashlib.sha256(original).hexdigest()},ensure_ascii=False))
