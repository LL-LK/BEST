"""编辑已有论文稿的新版本，不覆盖原交付或冻结实验材料。"""
import argparse
import copy
import json
from pathlib import Path
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.shared import Cm


ROOT = Path(__file__).resolve().parents[2]
OLD = ROOT / 'docs/results/2026-10-11/e3-e6-paper'


def append_table(doc, template, rows, widths):
    element = copy.deepcopy(template._tbl)
    doc._body._body.insert(len(doc._body._body)-1, element)
    table = Table(element, doc._body)
    header, body = copy.deepcopy(table.rows[0]._tr), copy.deepcopy(table.rows[1]._tr)
    for row in list(table.rows):
        table._tbl.remove(row._tr)
    for index, values in enumerate(rows):
        table._tbl.append(copy.deepcopy(header if index == 0 else body))
        cells = table.rows[-1].cells
        for cell, text, width in zip(cells, values, widths):
            # Preserve formatting of the retained table cell's first run.
            paragraph = cell.paragraphs[0]
            original = copy.deepcopy(paragraph.runs[0]._r)
            for child in list(paragraph._p):
                if child.tag.endswith('}r'):
                    paragraph._p.remove(child)
            paragraph._p.append(original)
            paragraph.runs[0].text = str(text)
            cell.width = Cm(width)
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    return table


def revise(out):
    result = json.loads((out/'FINAL-RESULTS.json').read_text())
    if not result['collection_complete'] or not result['delivery_integrity_audit']['passed']:
        raise ValueError('未完成采集和交付审计，不能生成完成稿')
    doc = Document(OLD/'E3与E6-论文方法与结果.docx')
    tables = doc.tables
    metrics_template, usability_template = tables[0], tables[4]
    replacements = {
        'E3 的回答采集、助手定量评分及教师定性反馈整理已完成。E6 的 R0、R1、R3 计划槽位均已结束，但包含真实失败，R2 跨日复测和 U7 主观评分尚未完成。现有数据可用于论文的探索性结果和失败分析；原协议的独立评阅与全部实验验收条件仍未满足。':
        'E3 的回答采集、助手定量评分及教师定性反馈整理已完成。E6 原批次的 R0、R1、R3 及充值后的独立 R3 补测均已结束，真实失败纳入统计。研究者于 10 月 11 日取消难以实施的主观 U7，当前方案仅报告客观可用性 U1–U6。R2 跨日复测仍待 10 月 14 日执行。现有数据可用于论文的探索性结果和失败分析；E3 独立定量评阅与原协议的全部验收条件仍未满足。',
        'E3 实际主回答采集时间为 2026 年 10 月 6 日 02:43:28 至 03:45:53（上海时间），E6 的 R0、R1、R3 实际执行日期均为 10 月 7 日。10 月 11 日进行离线收尾核验，没有新增计费模型调用。':
        'E3 实际主回答采集时间为 2026 年 10 月 6 日 02:43:28 至 03:45:53（上海时间），E6 原批次的 R0、R1、R3 于 10 月 7 日执行。10 月 11 日先完成原数据的离线复核，随后在更新认证密钥和充值后开始独立 R3 全量补测。补测调用真实生成模型 API，工具数据使用原冻结回包。',
        'E6 采集覆盖与一致性结果': 'E6 原批次采集覆盖与一致性结果',
        'E6 效率与可用性结果': 'E6 原批次效率与可用性结果',
        '原 12 项反证夹具全部检出，metascore = 1.0。冻结快照有 274 个请求键，619 个来源指针重新核验均通过；所有正式轨迹的冻结输入一致，没有密钥命中。原协议门禁仍为 FAIL，U7 与综合分保持缺失。':
        '原 12 项反证夹具全部检出，metascore = 1.0。冻结快照有 274 个请求键，619 个来源指针重新核验均通过；原正式轨迹的冻结输入一致，没有密钥命中。当前修订方案取消 U7，表中保留客观 U1–U6。原协议门禁仍为 FAIL，不据修订后指标补算原七项综合分。',
        '若补足额度后继续 R3，建议建立单独的补测批次，保留原 69 次余额失败和全部初始成功率。补测结果单独报告，不能把多次尝试中的成功答案择优合并成原批满分。主观 U7 需实际使用者独立评阅；E3 的一致性需完整独立量表和反证核验，两者不能由当前助手自评分补出。':
        '充值后的 R3 已另建全量补测批次，使用原 32 题各五次，结果见下一节。原 69 次余额失败及全部初始有效记录率保留，补测不择优替换原答案。U7 经研究者明确取消，不再作为 E6 待完成项；该修改发生在 R3 补测期间，此前没有任何 U7 评分。E3 一致性仍需要完整独立量表与反证核验，不能由当前助手补评分得到。',
        '本次复核通过 306 项现有及新增测试，pip check 无依赖冲突；E3 全部原始计数与 E6 各路统计均精确复现。复核包包含主表、分母、逐回答计数、逐题一致性、教师反馈修订报告、失败与效率明细、原件哈希，以及复跑命令。原始证据归档保存在仓库 docs/results/2026-10-07 两个实验目录中，来源与哈希见 SOURCE-MANIFEST.json。':
        '原数据离线复核精确复现 E3 全部计数与 E6 各路统计，环境检查无依赖冲突。补测流程加入后，现有及新增流程共 314 项测试通过。原复核包保留主表、分母、逐回答计数、逐题一致性、教师反馈及原件哈希；本次另交付实际补测轨迹、失败明细和可离线核验的快照来源，不覆盖原归档。',
    }
    observed = set()
    for p in doc.paragraphs:
        if p.text in replacements:
            observed.add(p.text)
            p.text = replacements[p.text]
    if observed != set(replacements):
        raise ValueError('历史稿结构改变，停止自动编辑')
    for row in list(usability_template.rows):
        if row.cells[0].text.startswith('U7'):
            usability_template._tbl.remove(row._tr)
    doc.paragraphs[1].text = '原批次与充值后补测分列  U7 已取消'
    heading = doc.add_paragraph('E6 充值后独立补测结果', 'Heading 1')
    heading.paragraph_format.page_break_before = True
    score = result['consistency']
    audit = result['collection_audit']
    doc.add_paragraph(
        f"本批次为事后全量补测，编号 {result['batch']}。32 题各五次，共 160 个计划槽位；沿用原模型、参数、系统提示、工具实现及冻结实际回包。每个槽位只有一次执行尝试，调用失败及输出缺失均保留，不重新尝试至成功。全部槽位已结束，有效记录 {score['records']}/160（{score['records']/160:.2%}），失败 {audit['failure_attempts']} 次，完整五次有效题数 {score['coverage']['tasks_with_5_valid']}/32。")
    doc.add_paragraph('以下一致性按有效记录计算，题号为重抽单位，bootstrap 10000 次。少于两次有效重复的题目不估计稳定性。有效记录包含符合格式的澄清与拒绝，不等于回答满足科学要求或任务一次可用。')
    rows = [['指标', '点估计', '95% bootstrap 区间', '可估计题数']]
    for name in ('TSC','PAC_call','RDC','FCC'):
        item = score['aggregate'][name]
        ci = item['bootstrap_95']
        rows.append([name, f"{item['point']:.4f}" if item['point'] is not None else '缺失',
                     f'[{ci[0]:.4f}, {ci[1]:.4f}]' if ci else '缺失', item['task_units']])
    append_table(doc, metrics_template, rows, [3.2,3.2,7.8,3.6])
    doc.add_paragraph('TSC 是工具序列一致性，PAC_call 是调用参数一致性，RDC 是工具回包指纹一致性，FCC 是核心结论一致性。FCC 采用当前助手及本地语义规则，评阅非独立，不代表全文科学正确。服务商权重版本不可核验。原批次与本次均单列，不将它们合并解释为同一时点的对照试验。')
    ef = result['efficiency']['EF7']
    doc.add_paragraph(
        f"本批次余额不足失败 {result['generation_api_failures']['insufficient_balance_attempts']} 次。全部尝试的已知输入 token 合计 {ef['tokens_in_observed_sum']:,}，输出 token 合计 {ef['tokens_out_observed_sum']:,}。缺失 usage 保持未知，货币费用未核验，不能作为完整账单。工具回放没有新增 BEST wire 请求；生成模型调用实际计入轨迹。")
    if (out/'OPERATIONAL-TIMING-NOTE.json').exists():
        note=json.loads((out/'OPERATIONAL-TIMING-NOTE.json').read_text())
        doc.add_paragraph(
            f"REPRO-22 第二次重复出现网络超时与分块响应中断，单调计时 {note['monotonic_wall_s']:.3f} 秒，日历跨度 {note['calendar_span_s']:.3f} 秒。差额原因尚未确认，不归为模型推理时间。随后临时防止采集进程空闲休眠；这一运行条件变化保留记录，未改变输入、生成参数或重试已结束槽位。")
    doc.add_paragraph('客观可用性继续使用各项实际适用分母，失败尝试也纳入相关分母。U7 已取消，没有赋零分、助手代评或补算原综合分。')
    rows = [['指标', '点估计', '实际适用分母']]
    labels = {'U1':'一次可用率','U2':'错误信息要素率','U3':'澄清要素率','U4':'溯源标注覆盖','U5':'步骤比值','U6':'降级标签率'}
    for name,label in labels.items():
        item=result['usability']['summary'][name]
        rows.append([f'{name} {label}', f"{item['point']:.4f}" if item['point'] is not None else '缺失', item['observations']])
    append_table(doc, usability_template, rows, [8.0,4.9,4.9])
    doc.add_paragraph('全部失败及其错误原件保留，原批次文件、冻结输入和快照来源核验通过。一条 NaN 工具回包序列化失败的轨迹存在工具结果及成本对账缺段，保持无效；不能称全部轨迹完整。收尾时配置已无 API 密钥，统计仅离线恢复。凭据结构与模式扫描通过，按实际密钥进行的精确扫描不可用，不能标为已通过。采集完成、输出有效和科学表现合格分别报告。温度、题目标记、材料及覆盖限制继续披露；D+7 R2 尚未执行，原完整协议未通过。')
    path=out/'E3与E6-论文方法与结果-修订版.docx'
    doc.save(path)
    # Export matching readable text in document order.
    blocks=[]
    for child in doc._body._body:
        if child.tag.endswith('}p'):
            p=Paragraph(child,doc._body)
            if p.text:
                prefix='# ' if p.style.name=='Title' else '## ' if p.style.name=='Heading 1' else ''
                blocks.append(prefix+p.text)
        elif child.tag.endswith('}tbl'):
            table=Table(child,doc._body)
            rows=[[c.text.replace('\n',' ') for c in r.cells] for r in table.rows]
            blocks.append('\n'.join(['| '+' | '.join(rows[0])+' |', '| '+' | '.join(['---']*len(rows[0]))+' |']+
                                    ['| '+' | '.join(row)+' |' for row in rows[1:]]))
    (out/'E3与E6-论文方法与结果-修订版.md').write_text('\n\n'.join(blocks)+'\n')
    print(path)


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--out',type=Path,required=True)
    revise(parser.parse_args().out)
