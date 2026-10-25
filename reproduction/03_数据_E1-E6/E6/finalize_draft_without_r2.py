from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import json, shutil
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT=Path(__file__).resolve().parents[2]
B='E6-R3-supplement-20261011-1791689271010679000'
OUT=ROOT/'docs/results/2026-10-11/e6-r3-supplement'/B
AT=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat()
HIST=ROOT/'temp/e6-formal/publication-drafts'/('before-R2-cancellation-'+datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%H%M%S'))
shutil.copytree(OUT,HIST)
def dump(p,d): p.write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
amend={'at':AT,'revision':'E6-batch-repetition-consistency-U1-U6-no-R2','researcher_instruction':'为什么要复测，可以删除吗，我们今天就要初稿','R2_status':'cancelled_before_execution','R2_pending':False,'U7_status':'cancelled','retained_lanes':['R0','R1','R3-original','R3-supplement'],'retained_usability_metrics':['U1','U2','U3','U4','U5','U6'],'current_scope':'各批次内重复运行一致性、冻结快照回放、效率及客观可用性；原批次与事后补测分列','unmeasured_constructs':['D+7跨日稳定性','服务商权重及在线后端跨日漂移'],'remaining_E6_collection':None,'future_API_required_for_current_scope':False,'timing':'原R1/R3与R3补测结果已知后的事后范围修订；未执行R2','effect_on_raw_records':'none; 原协议、原始回答、失败和快照归档保留','original_protocol_pass':False,'historical_publication_snapshot':str(HIST)}
dump(OUT/'PROTOCOL-AMENDMENT-R2-CANCELLATION.json',amend)
dump(ROOT/'temp/e6-formal/CURRENT-PROTOCOL-AMENDMENT.json',amend)
for name in ['FINAL-RESULTS.json','SUPPLEMENT-PROTOCOL.json']:
 p=OUT/name; d=json.loads(p.read_text()); d['current_protocol_amendment']=amend
 d['R2']={'status':'cancelled_before_execution','pending':False,'collected_records':0,'reason':'研究者取消本次初稿的跨日实验范围','cross_day_stability_verified':False,'original_plan_archived':True}
 d['current_scope_execution_closed']=True
 if name=='FINAL-RESULTS.json':
  d['limitations'].append('R2 跨日复测于结果已知后取消；本研究未验证 D+7 稳定性或在线跨日漂移，不能据本批次一致率推论。')
  d['revised_scope_scientific_pass']=False
 dump(p,d)
report=OUT/'REPORT.md'; s=report.read_text(); old='原 R1 已结束 160 个计划槽位，其格式、工具及科学失败不重试至成功。D+7 R2 仍须在 **2026-10-14（Asia/Shanghai）** 使用原正式批次运行；本次 R3 补测不改变该日期。'; assert old in s
s=s.replace(old,'原 R1 已结束 160 个计划槽位，其格式、工具及科学失败不重试至成功。研究者于 2026 年 10 月 11 日在既有结果已知后取消 R2 D+7 跨日复测，当前 E6 采集范围已结束，无待执行模型调用。本研究仅分析各批次内重复运行一致性和冻结快照回放，未验证跨日稳定性或在线后端漂移。原批次与事后补测单独报告，不能作为控制其他条件的跨日比较。')
report.write_text(s)
p=OUT/'论文补充段落.md'; s=p.read_text(); s=s.replace('D+7 在线 R2 仍须在 2026 年 10 月 14 日执行。','研究者于 2026 年 10 月 11 日在结果已知后取消 R2 跨日复测。当前 E6 采集范围已结束，只报告各批次内重复一致性与冻结回包重放，不声称验证 D+7 稳定性或在线后端跨日漂移。'); p.write_text(s)
source=OUT/'E3与E6-论文方法与结果-修订版.docx'; doc=Document(source)
changes=0
for p in doc.paragraphs:
 s=p.text
 if 'R2 跨日复测仍待 10 月 14 日执行。' in s:
  s=s.replace('R2 跨日复测仍待 10 月 14 日执行。','研究者于同日取消 R2 跨日复测，当前 E6 采集范围已结束；结果仅支持各批次内重复一致性分析，不验证跨日稳定性。')
 if s.startswith('R2 的 D+7 日期为'):
  s='原方案的 R2 用于比较 R1 与七天后的真实在线运行，以评估跨日稳定性及后端漂移。研究者于 2026 年 10 月 11 日在原批次和 R3 补测结果已知后取消该路线，R2 未执行，也不列为当前实验待完成项。这是事后范围修订，原协议仍保留。论文仅分析各批次内五次重复的一致性及冻结快照回放；原 R3 与充值后补测分列，不能将两者当作控制其他条件的跨日验证。'
 s=s.replace('D+7 R2 尚未执行，原完整协议未通过。','R2 跨日复测已取消，跨日稳定性未验证；原完整协议未通过，当前范围内的已知科学失败仍保留。')
 if s!=p.text: p.text=s; changes+=1
assert changes==3,changes
removed=0
for table in doc.tables:
 for row in list(table.rows):
  if row.cells[0].text.strip()=='R2': table._tbl.remove(row._tr); removed+=1
assert removed==1
subtitle=doc.paragraphs[1]; subtitle.text='原批次与充值后补测分列  R2 与 U7 已取消'
path=OUT/'E3与E6-论文初稿.docx'; doc.save(path)
blocks=[]
for child in doc._body._body:
 if child.tag.endswith('}p'):
  p=Paragraph(child,doc._body)
  if p.text: blocks.append(('# ' if p.style.name=='Title' else '## ' if p.style.name=='Heading 1' else '')+p.text)
 elif child.tag.endswith('}tbl'):
  t=Table(child,doc._body); rows=[[c.text.replace('\n',' ') for c in r.cells] for r in t.rows]
  blocks.append('\n'.join(['| '+' | '.join(rows[0])+' |','| '+' | '.join(['---']*len(rows[0]))+' |']+['| '+' | '.join(row)+' |' for row in rows[1:]]))
path.with_suffix('.md').write_text('\n\n'.join(blocks)+'\n')
(OUT/'REMAINING-WORK.md').write_text('''# 当前实验范围与完成边界

研究者于 2026 年 10 月 11 日在已有结果之后取消 R2 D+7 跨日复测；U7 主观评分已先行取消。E6 当前约定的采集范围已结束，没有待执行模型调用或必须填写的 API key，也没有定时任务。

当前报告保留 R0、R1、原 R3 和独立 R3 补测，分析各批次内重复运行一致性、冻结快照回放、效率及客观 U1–U6。不声称验证了跨日稳定性、模型权重漂移或在线后端跨日变化。新补测 160 次尝试中 133 份有效、27 份失败；原失败和历史协议不删除。

取消未执行路线不改变既有失败或补算原完整协议通过。温度、任务标记、覆盖及材料错误继续披露。E3 独立定量量表不足和 M6 缺失也仍保留；教师定性反馈可用于讨论，不能代替完整独立定量评分。本初稿可用于论文写作，相关结论需采用上述边界。
''')
print(json.dumps({'new_document':str(path),'changed_paragraphs':changes,'R2_table_rows_removed':removed,'historical_snapshot':str(HIST)},ensure_ascii=False))
