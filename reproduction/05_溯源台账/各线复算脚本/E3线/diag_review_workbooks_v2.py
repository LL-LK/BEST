# -*- coding: utf-8 -*-
"""E3 匿名审核工作簿损坏诊断 v2：跨表引用完整性 + 缓存错误值（只读）。"""
import json, re, zipfile, collections, sys, io
from pathlib import Path
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

ROOT = Path(r'<WORKSPACE>\04_实验数据\E3_科学忠实性\实验产物')
OUT = Path(r'<WORKSPACE>\11_投稿包\_work\E1E2E3\_work_diag')
OUT.mkdir(parents=True, exist_ok=True)

TARGETS = [
    ROOT / '04_匿名审核材料' / 'E3-全量匿名审核表.xlsx',
    ROOT / '04_匿名审核材料' / 'E3-匿名审核简表.xlsx',
    ROOT / '02_教师反馈' / 'E3-老师审核表-修复版.xlsx',
]

report = {}
for p in TARGETS:
    info = {'path': str(p), 'exists': p.exists()}
    if not p.exists():
        report[p.name] = info
        continue
    info['size'] = p.stat().st_size
    with zipfile.ZipFile(p) as z:
        names = z.namelist()
        ref_parts, err_cells = {}, collections.Counter()
        for n in names:
            if n.endswith('.xml'):
                raw = z.read(n).decode('utf8', 'replace')
                c = raw.count('#REF!')
                if c:
                    ref_parts[n] = c
                for m in re.finditer(r'<c\b[^>]*t="e"[^>]*>.*?</c>', raw, re.S):
                    err_cells[n] += 1
        info['ref_error_text_by_xml'] = ref_parts
        info['ref_error_text_total'] = sum(ref_parts.values())
        info['cached_error_cells_by_xml'] = dict(err_cells)
        info['cached_error_cells_total'] = sum(err_cells.values())
        info['calcChain_present'] = 'xl/calcChain.xml' in names
        info['externalLinks'] = [n for n in names if 'externalLink' in n]

    wb = openpyxl.load_workbook(p, data_only=False)
    sheetnames = wb.sheetnames
    info['sheetnames'] = sheetnames

    # 跨表引用检查
    dangling = collections.Counter()
    formula_by_sheet = {}
    cross_refs = set()
    for ws in wb.worksheets:
        fs = []
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('='):
                    fs.append((c.coordinate, c.value))
                    for m in re.finditer(r"'([^']+)'!", c.value):
                        cross_refs.add(m.group(1))
                        if m.group(1) not in sheetnames:
                            dangling[m.group(1)] += 1
                    for m in re.finditer(r"(?<!')\b([\u4e00-\u9fffA-Za-z0-9_]+)!", c.value):
                        ref = m.group(1)
                        if ref not in sheetnames and ref not in ('SUM', 'COUNT', 'IF', 'COUNTIFS'):
                            dangling[ref] += 1
        formula_by_sheet[ws.title] = {'count': len(fs), 'sample': fs[:3]}
    info['formula_by_sheet'] = formula_by_sheet
    info['cross_sheet_refs_used'] = sorted(cross_refs)
    info['dangling_sheet_refs'] = dict(dangling)
    info['dangling_total'] = sum(dangling.values())

    # 缓存值(重算结果)中的错误
    wb2 = openpyxl.load_workbook(p, data_only=True)
    cached_err = collections.Counter()
    none_formula = 0
    for ws in wb2.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('#'):
                    cached_err[c.value] += 1
    info['data_only_error_values'] = dict(cached_err)

    # 首行表头（判读结构）
    wb3 = openpyxl.load_workbook(p, data_only=False)
    for ws in wb3.worksheets:
        hdr = []
        for c in ws[1]:
            if c.value is not None:
                hdr.append(f'{c.coordinate}={c.value}')
        if hdr:
            info.setdefault('headers', {})[ws.title] = hdr[:30]
    report[p.name] = info

(OUT / 'xlsx_ref_diagnosis_v2.json').write_text(
    json.dumps(report, ensure_ascii=False, indent=2), encoding='utf8')

for name, info in report.items():
    print('=' * 78)
    print(name, '| exists=', info.get('exists'), '| size=', info.get('size'))
    if not info.get('exists'):
        continue
    print('  sheets            :', info['sheetnames'])
    print('  "#REF!" text in XML:', info['ref_error_text_total'], info['ref_error_text_by_xml'])
    print('  cached error cells:', info['cached_error_cells_total'], info['cached_error_cells_by_xml'])
    print('  data_only errors  :', info['data_only_error_values'])
    print('  calcChain present :', info['calcChain_present'])
    print('  externalLinks     :', info['externalLinks'])
    print('  cross-sheet refs  :', info['cross_sheet_refs_used'])
    print('  DANGLING refs     :', info['dangling_sheet_refs'], '=> total', info['dangling_total'])
    for t, v in info['formula_by_sheet'].items():
        print(f"   [{t}] formulas={v['count']}")
