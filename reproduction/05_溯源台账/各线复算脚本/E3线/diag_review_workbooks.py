# -*- coding: utf-8 -*-
"""诊断 E3 匿名审核工作簿的 #REF! 损坏范围（只读，不写原始文件）。"""
import json, re, zipfile, collections
from pathlib import Path
import openpyxl

SRC = Path(r'<WORKSPACE>\04_实验数据\E3_科学忠实性\实验产物\04_匿名审核材料')
OUT = Path(r'<WORKSPACE>\11_投稿包\_work\E1E2E3\_work_diag')
OUT.mkdir(parents=True, exist_ok=True)

report = {}
for name in ['E3-全量匿名审核表.xlsx', 'E3-匿名审核简表.xlsx']:
    p = SRC / name
    info = {'path': str(p), 'size': p.stat().st_size}
    # 1) 原始 XML 层面统计
    with zipfile.ZipFile(p) as z:
        names = z.namelist()
        info['zip_entries'] = len(names)
        ref_cnt = collections.Counter()
        formula_cnt = 0
        for n in names:
            if n.endswith('.xml'):
                raw = z.read(n).decode('utf8', 'replace')
                c = raw.count('#REF!')
                if c:
                    ref_cnt[n] = c
                formula_cnt += len(re.findall(r'<f[ >]', raw))
        info['ref_error_by_xml'] = dict(ref_cnt)
        info['ref_error_total'] = sum(ref_cnt.values())
        info['formula_tags_total'] = formula_cnt
        # 定义名称
        for n in names:
            if n.endswith('workbook.xml'):
                wbxml = z.read(n).decode('utf8', 'replace')
                dn = re.findall(r'<definedName[^>]*name="([^"]+)"[^>]*>([^<]*)</definedName>', wbxml)
                info['defined_names'] = dn
                info['calcChain_present'] = 'xl/calcChain.xml' in names
                m = re.search(r'<calcPr[^>]*/>', wbxml)
                info['calcPr'] = m.group(0) if m else None
    # 2) openpyxl 读取每个 sheet 的公式/错误
    wb = openpyxl.load_workbook(p, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        formulas, errors, nonempty = [], collections.Counter(), 0
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                nonempty += 1
                if isinstance(c.value, str) and c.value.startswith('='):
                    formulas.append((c.coordinate, c.value))
                if isinstance(c.value, str) and c.value.startswith('#'):
                    errors[c.value] += 1
        sheets.append({
            'title': ws.title, 'max_row': ws.max_row, 'max_column': ws.max_column,
            'nonempty_cells': nonempty, 'formula_count': len(formulas),
            'formula_samples': formulas[:5],
            'literal_error_values': dict(errors),
        })
    info['sheets'] = sheets
    report[name] = info

(OUT / 'xlsx_ref_diagnosis.json').write_text(
    json.dumps(report, ensure_ascii=False, indent=2), encoding='utf8')

for name, info in report.items():
    print('=' * 70)
    print(name, 'size=', info['size'])
    print('  #REF! total in XML :', info['ref_error_total'])
    print('  by xml part         :', info['ref_error_by_xml'])
    print('  <f> formula tags    :', info['formula_tags_total'])
    print('  calcChain present   :', info.get('calcChain_present'))
    print('  calcPr              :', info.get('calcPr'))
    print('  defined_names       :', info.get('defined_names'))
    for s in info['sheets']:
        print(f"  - sheet {s['title']!r}: rows={s['max_row']} cols={s['max_column']} "
              f"nonempty={s['nonempty_cells']} formulas={s['formula_count']} "
              f"literal_errors={s['literal_error_values']}")
        for co, f in s['formula_samples']:
            print(f"      {co}: {f[:160]}")
