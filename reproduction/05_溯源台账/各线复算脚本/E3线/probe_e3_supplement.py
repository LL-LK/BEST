# -*- coding: utf-8 -*-
"""E3 补充核对：红线计数 / 在线辅测 / 归档 source 字段 / 匿名审核表结构（只读）。"""
from pathlib import Path
from collections import Counter, defaultdict
import json, csv, tarfile, io, sys, hashlib
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

E3 = Path(r'<WORKSPACE>\04_实验数据\E3_科学忠实性\实验产物')
ARC = E3 / '05_原始实验归档'
CUR = E3 / '01_当前结果'
DIAG = Path(r'<WORKSPACE>\11_投稿包\_work\E1E2E3\_work_diag')
DIAG.mkdir(parents=True, exist_ok=True)
out = {}

bundle = json.loads((ARC / 'bundle.json').read_text('utf-8-sig'))
print('bundle top-level keys:', list(bundle.keys()))
out['bundle_keys'] = list(bundle.keys())
for k in bundle:
    v = bundle[k]
    if isinstance(v, list):
        print(f'  {k}: list n={len(v)}')
    elif isinstance(v, dict):
        print(f'  {k}: dict keys={list(v.keys())[:12]}')
    else:
        print(f'  {k}: {type(v).__name__} len={len(str(v))}')

answers, scores = {}, {}
with tarfile.open(ARC / 'original-evidence.tar.gz', 'r:gz') as z:
    for n in z.getnames():
        if n.endswith('.json') and ('/answers/' in n or '/scores/' in n):
            a = json.loads(z.extractfile(n).read())
            (answers if '/answers/' in n else scores)[a['key']] = a

# ---------- 1. 红线 / fatal / 禁用项 ----------
red = defaultdict(list)
for key, s in scores.items():
    a = answers.get(key)
    if a is None:
        continue
    g = a['group']
    rl = s['counts'].get('redlines') or []
    if rl:
        red[g].append({'key': key, 'n_redlines': len(rl), 'redlines': rl})
    for cl in s['score'].get('redlines') or []:
        pass
print('\n--- 红线 ---')
for g, v in red.items():
    print(f'  {g}: 含红线回答数={len(v)} 红线条数合计={sum(x["n_redlines"] for x in v)}')
    for x in v[:3]:
        print('   ', x['key'], x['n_redlines'], json.dumps(x['redlines'], ensure_ascii=False)[:220])
out['redlines'] = {g: {'answers_with_redline': len(v), 'redline_total': sum(x['n_redlines'] for x in v),
                        'keys': [x['key'] for x in v]} for g, v in red.items()}

# ---------- 2. 归档 claim 字段完整性（supported + unsupported） ----------
fieldcnt = Counter()
src_nonempty = Counter()
cls = Counter()
for key, s in scores.items():
    for cl in s['score']['claims']:
        fieldcnt['total_claims'] += 1
        cls[cl['classification']] += 1
        for f in ('quote', 'reason', 'source'):
            if f in cl:
                fieldcnt['has_' + f] += 1
            if (cl.get(f) or '').strip():
                fieldcnt['nonempty_' + f] += 1
        if cl['classification'] == 'unsupported':
            src_nonempty['unsupported'] += 1 if (cl.get('source') or '').strip() else 0
            src_nonempty['unsupported_total'] += 1
        else:
            src_nonempty['supported'] += 1 if (cl.get('source') or '').strip() else 0
            src_nonempty['supported_total'] += 1
print('\n--- 归档 claim 字段完整性 ---')
print('  ', dict(fieldcnt))
print('   classification:', dict(cls))
print('   source 非空: unsupported %d/%d ; supported %d/%d' % (
    src_nonempty['unsupported'], src_nonempty['unsupported_total'],
    src_nonempty['supported'], src_nonempty['supported_total']))
out['claim_field_completeness'] = {'counts': dict(fieldcnt), 'classification': dict(cls),
                                   'source_nonempty': dict(src_nonempty)}

# ---------- 3. numbers / 理由最小长度 ----------
short_reason = []
for key, s in scores.items():
    for cl in s['score']['claims']:
        if len((cl.get('reason') or '').strip()) < 8:
            short_reason.append((key, cl.get('quote', '')[:30], cl.get('reason', '')))
print(f'   理由 <8 字的 claim: {len(short_reason)}')
out['short_reasons'] = short_reason[:20]

# ---------- 4. 在线辅测 ----------
anc = json.loads((CUR / 'E3-online-ancillary.json').read_text('utf-8-sig'))
out['online_ancillary'] = anc
print('\n--- 在线辅测 E3-online-ancillary.json 顶层 ---')
for k, v in anc.items():
    print('   ', k, ':', json.dumps(v, ensure_ascii=False)[:300])

# ---------- 5. 四条核验 ----------
fv = json.loads((E3 / '03_四条核验' / 'four-entry-verification.json').read_text('utf-8-sig'))
out['four_entry_verification'] = fv
print('\n--- 四条核验 top keys ---', list(fv.keys()))

# ---------- 6. 匿名审核表结构 ----------
wb_path = E3 / '04_匿名审核材料' / 'E3-全量匿名审核表.xlsx'
wb = openpyxl.load_workbook(wb_path, data_only=False)
print('\n--- 全量匿名审核表 sheets ---', wb.sheetnames)
sheet_info = {}
for ws in wb.worksheets:
    rows = []
    for r in ws.iter_rows(min_row=1, max_row=8):
        vals = [(c.coordinate, c.value) for c in r if c.value is not None]
        rows.append(vals)
    sheet_info[ws.title] = {'max_row': ws.max_row, 'max_col': ws.max_column, 'head_rows': rows}
    print(f'\n  [{ws.title}] rows={ws.max_row} cols={ws.max_column}')
    for rr in rows[:6]:
        print('     ', [(co, (str(v)[:40] if not isinstance(v, (int, float)) else v)) for co, v in rr][:14])
out['review_workbook_structure'] = sheet_info

( DIAG / 'e3_supplement_v1.json').write_text(json.dumps(out, ensure_ascii=False, indent=2, default=str), encoding='utf8')
print('\n->', DIAG / 'e3_supplement_v1.json')
