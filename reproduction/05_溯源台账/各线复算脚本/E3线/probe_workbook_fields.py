# -*- coding: utf-8 -*-
"""核对全量匿名审核表的条目来源、超链接、数据校验、参考材料覆盖，以及 samples/ 源是否仍存在。"""
from pathlib import Path
import json, io, sys, collections
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
E3 = Path(r'<WORKSPACE>\04_实验数据\E3_科学忠实性\实验产物')
DIAG = Path(r'<WORKSPACE>\11_投稿包\_work\E1E2E3\_work_diag')
out = {}

print('=== expert-package-status.json ===')
print((E3 / '05_原始实验归档' / 'expert-package-status.json').read_text('utf-8-sig'))
print('=== judge-selftest.json ===')
print((E3 / '05_原始实验归档' / 'judge-selftest.json').read_text('utf-8-sig'))
print('=== agreement.json ===')
print((E3 / '05_原始实验归档' / 'agreement.json').read_text('utf-8-sig'))

# samples/ 目录是否仍在
for base in [E3 / '05_原始实验归档', E3 / '04_匿名审核材料', E3, Path(r'<WORKSPACE>\04_实验数据')]:
    hits = list(base.rglob('samples')) if base.exists() else []
    if hits:
        print(f'  samples/ found under {base}: {hits[:5]}')
out['samples_dirs_found'] = [str(h) for b in [E3, Path(r'<WORKSPACE>\04_实验数据')] if b.exists()
                             for h in list(b.rglob('samples'))[:20]]

wb = openpyxl.load_workbook(E3 / '04_匿名审核材料' / 'E3-全量匿名审核表.xlsx', data_only=False)

print('\n=== 逐项判定：类别分布 & 样本1全部行 ===')
ws = wb['逐项判定']
cats = collections.Counter()
sample1 = []
for r in range(7, ws.max_row + 1):
    b = ws.cell(row=r, column=2).value
    c = ws.cell(row=r, column=3).value
    cats[c] += 1
    if b == '943d634b0934e17f':
        sample1.append([ws.cell(row=r, column=i).value for i in range(1, 14)])
print(' 类别:', dict(cats))
for row in sample1:
    print('  ', [str(x)[:70] if x is not None else None for x in row])
out['per_item_categories'] = dict(cats)
out['per_item_sample1'] = sample1

print('\n=== 数值登记 前9行 / 主张登记 前9行 ===')
for sn in ['数值登记', '主张登记']:
    w = wb[sn]
    print(f' [{sn}] max_row={w.max_row}')
    for r in range(7, min(w.max_row, 15) + 1):
        vals = [(c.coordinate, str(c.value)[:60]) for c in w[r] if c.value is not None]
        print('   ', vals)

print('\n=== 参考材料 类别分布 ===')
ws = wb['参考材料']
rc = collections.Counter()
for r in range(7, ws.max_row + 1):
    rc[ws.cell(row=r, column=1).value] += 1
print(' ', dict(rc))
for r in range(7, 16):
    print('   ', [(c.coordinate, str(c.value)[:80]) for c in ws[r] if c.value is not None])

print('\n=== 超链接 / 数据校验 ===')
for name in wb.sheetnames:
    w = wb[name]
    links = [(c.coordinate, c.hyperlink.target if c.hyperlink else None, c.value)
             for row in w.iter_rows() for c in row if c.hyperlink]
    dvs = []
    for dv in w.data_validations.dataValidation:
        dvs.append({'type': dv.type, 'formula1': str(dv.formula1)[:200], 'sqref': str(dv.sqref)[:200]})
    print(f' [{name}] hyperlinks={len(links)} datavalidations={len(dvs)}')
    if links[:2]:
        print('    e.g.', links[:2])
    for dv in dvs[:4]:
        print('    dv:', dv)
    out.setdefault('links', {})[name] = links[:5]
    out.setdefault('dvs', {})[name] = dvs

(DIAG / 'e3_workbook_probe.json').write_text(json.dumps(out, ensure_ascii=False, indent=2, default=str), encoding='utf8')
print('\n->', DIAG / 'e3_workbook_probe.json')
