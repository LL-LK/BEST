# -*- coding: utf-8 -*-
"""核对：内部超链接是否有效、219 匿名编号 ↔ 归档 answer_key 映射是否可闭合。"""
from pathlib import Path
from collections import Counter
import json, io, sys, tarfile
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
E3 = Path(r'<WORKSPACE>\04_实验数据\E3_科学忠实性\实验产物')
ARC = E3 / '05_原始实验归档'
DIAG = Path(r'<WORKSPACE>\11_投稿包\_work\E1E2E3\_work_diag')

wb = openpyxl.load_workbook(E3 / '04_匿名审核材料' / 'E3-全量匿名审核表.xlsx', data_only=False)
names = set(wb.sheetnames)
bad = []
tot = 0
for sn in wb.sheetnames:
    ws = wb[sn]
    for row in ws.iter_rows():
        for c in row:
            if c.hyperlink:
                tot += 1
                hl = c.hyperlink
                tgt = getattr(hl, 'target', None)
                loc = getattr(hl, 'location', None)
                if tgt is None and not loc:
                    bad.append((sn, c.coordinate, repr(hl)))
                elif loc and loc.split('!')[0].strip("'") not in names:
                    bad.append((sn, c.coordinate, 'bad-location:' + str(loc)))
print(f'内部/外部超链接总数={tot}  无效={len(bad)}')
for b in bad[:10]:
    print('   ', b)

# 映射
recs = json.loads((E3 / '02_教师反馈' / 'teacher-review-records.json').read_text('utf-8-sig'))
if isinstance(recs, dict):
    print('\nrecords top keys:', list(recs.keys()))
    sample_key = recs.get('records', recs)
else:
    sample_key = recs
print('records type:', type(sample_key), 'len:', len(sample_key))
if isinstance(sample_key, list):
    print('  rec0:', json.dumps(sample_key[0], ensure_ascii=False)[:400])
    keys = [r.get('answer_key') for r in sample_key]
    print('  有 answer_key 的条数:', sum(1 for k in keys if k))
    print('  answer_key 样例:', [k for k in keys if k][:6])
    print('  非真实答案(control)样例:', [k for k in keys if k and k.startswith('control')][:5])
    ids = [r.get('sample_id') for r in sample_key]
    print('  sample_id 唯一:', len(set(ids)) == len(ids), len(set(ids)))

answers = {}
with tarfile.open(ARC / 'original-evidence.tar.gz', 'r:gz') as z:
    for n in z.getnames():
        if n.endswith('.json') and '/answers/' in n:
            a = json.loads(z.extractfile(n).read())
            answers[a['key']] = a
print('\n归档 answer keys:', len(answers))
if isinstance(sample_key, list):
    mapped = {r.get('answer_key') for r in sample_key if r.get('answer_key')}
    print('  映射覆盖的归档键:', len(mapped & set(answers)), '/ 归档', len(answers))
    print('  未映射的归档键:', sorted(set(answers) - mapped)[:10])
    ctrl = {k for k in mapped if k.startswith('control')}
    print('  control 键:', ctrl)

# 审核总表的 219 匿名编号
ws = wb['审核总表']
sids = [ws.cell(row=r, column=2).value for r in range(7, 226)]
print('\n审核总表 219 匿名编号唯一:', len(set(sids)) == 219, len(set(sids)))
if isinstance(sample_key, list):
    rec_ids = {r.get('sample_id') for r in sample_key}
    print('  与 teacher-review-records sample_id 交集:', len(set(sids) & rec_ids))
    print('  记录表多余:', sorted(rec_ids - set(sids))[:10])

json.dump({'hyperlinks_total': tot, 'hyperlinks_bad': bad[:50], 'hyperlinks_bad_n': len(bad),
           'audit_sids_n': len(set(sids))},
          open(DIAG / 'e3_link_and_mapping_probe.json', 'w', encoding='utf8'), ensure_ascii=False, indent=2)
print('\n->', DIAG / 'e3_link_and_mapping_probe.json')
